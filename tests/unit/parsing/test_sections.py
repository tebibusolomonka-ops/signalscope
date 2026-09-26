import io
import json

import docx
import pytest

from sample_files import make_pdf
from signalscope.parsing.docx_document import DOCX_CONTENT_TYPE
from signalscope.parsing.registry import create_default_parser_registry
from signalscope.parsing.types import (
    SECTION_SEPARATOR,
    MetadataValue,
    ParsedDocument,
    ParsedSection,
    single_section,
)


def parse(data: bytes, content_type: str) -> ParsedDocument:
    return create_default_parser_registry().get(content_type).parse(data, content_type=content_type)


def summary(document: ParsedDocument) -> list[tuple[str, int, str, dict[str, object]]]:
    return [
        (section.kind, section.index, section.text, dict(section.metadata))
        for section in document.sections
    ]


def check_consistent(document: ParsedDocument) -> None:
    assert SECTION_SEPARATOR.join(section.text for section in document.sections) == document.text


def test_section_metadata_cannot_be_changed() -> None:
    metadata: dict[str, MetadataValue] = {"page_number": 1}
    section = ParsedSection(text="Hi", kind="page", index=0, metadata=metadata)
    metadata["page_number"] = 2

    assert section.metadata == {"page_number": 1}
    with pytest.raises(TypeError):
        section.metadata["page_number"] = 3  # type: ignore[index]


@pytest.mark.parametrize(
    ("text", "sections", "message"),
    [
        ("A", (ParsedSection(text="A", kind="page", index=1),), "indexes must count from 0"),
        ("", (ParsedSection(text="", kind="page", index=0),), "must not be empty"),
        (
            "A B",
            (
                ParsedSection(text="A", kind="page", index=0),
                ParsedSection(text="B", kind="page", index=1),
            ),
            "section texts joined",
        ),
    ],
)
def test_sections_must_match_the_text(
    text: str, sections: tuple[ParsedSection, ...], message: str
) -> None:
    with pytest.raises(ValueError, match=message):
        ParsedDocument(text=text, sections=sections)


def test_document_without_sections_is_allowed() -> None:
    assert ParsedDocument(text="Hello").sections == ()


def test_single_section() -> None:
    assert single_section("", "document") == ()
    assert single_section("Hello", "body") == (ParsedSection(text="Hello", kind="body", index=0),)


def test_plain_text_is_one_section() -> None:
    document = parse(b"First.\n\nSecond.", "text/plain")

    assert summary(document) == [("document", 0, "First.\n\nSecond.", {})]
    check_consistent(document)


def test_empty_text_has_no_sections() -> None:
    assert parse(b"   ", "text/plain").sections == ()


def test_json_is_one_section() -> None:
    document = parse(json.dumps({"title": "Budget"}).encode(), "application/json")

    assert summary(document) == [("document", 0, "title: Budget", {})]


def test_html_body_is_one_section() -> None:
    document = parse(b"<html><body><p>Salut</p></body></html>", "text/html")

    assert summary(document) == [("body", 0, "Salut", {})]


def test_pdf_pages_are_sections() -> None:
    document = parse(make_pdf(["First page", "Second page\nwith two lines"]), "application/pdf")

    assert summary(document) == [
        ("page", 0, "First page", {"page_number": 1}),
        ("page", 1, "Second page\nwith two lines", {"page_number": 2}),
    ]
    check_consistent(document)


def test_pdf_pages_without_text_keep_the_real_page_numbers() -> None:
    document = parse(make_pdf(["One", None, "Three"]), "application/pdf")

    assert [section.metadata["page_number"] for section in document.sections] == [1, 3]
    assert [section.index for section in document.sections] == [0, 1]
    check_consistent(document)


def test_pdf_without_text_has_no_sections() -> None:
    assert parse(make_pdf([None, None]), "application/pdf").sections == ()


def docx_with_headings() -> bytes:
    document = docx.Document()
    document.add_paragraph("Summary before any heading.")
    document.add_heading("Background", level=1)
    document.add_paragraph("Why the work started.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Year"
    table.cell(0, 1).text = "2026"
    document.add_heading("Results", level=2)
    document.add_paragraph("What was found.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_is_split_at_headings() -> None:
    document = parse(docx_with_headings(), DOCX_CONTENT_TYPE)

    assert summary(document) == [
        ("section", 0, "Summary before any heading.", {}),
        (
            "section",
            1,
            "Background\n\nWhy the work started.\n\nYear\n\n2026",
            {"heading": "Background"},
        ),
        ("section", 2, "Results\n\nWhat was found.", {"heading": "Results"}),
    ]
    check_consistent(document)


def test_docx_without_headings_is_one_section() -> None:
    document = docx.Document()
    for text in ["One.", "Two."]:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert summary(parsed) == [("section", 0, "One.\n\nTwo.", {})]


def test_empty_docx_has_no_sections() -> None:
    buffer = io.BytesIO()
    docx.Document().save(buffer)

    assert parse(buffer.getvalue(), DOCX_CONTENT_TYPE).sections == ()
