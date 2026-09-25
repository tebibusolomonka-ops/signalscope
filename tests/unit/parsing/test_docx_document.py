import io
import struct
import zipfile
import zlib

import docx
import pytest

from signalscope.parsing import docx_document
from signalscope.parsing.docx_document import DOCX_CONTENT_TYPE, DocxDocumentParser
from signalscope.parsing.types import DocumentParsingError, ParsedDocument


def save(document: docx.document.Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def new_document(title: str = "", author: str = "") -> docx.document.Document:
    document = docx.Document()
    document.core_properties.title = title
    document.core_properties.author = author
    return document


def parse(data: bytes) -> ParsedDocument:
    return DocxDocumentParser().parse(data, content_type=DOCX_CONTENT_TYPE)


def test_paragraphs() -> None:
    document = new_document()
    document.add_heading("Annual report", level=1)
    document.add_paragraph("First paragraph.")
    document.add_paragraph("")
    document.add_paragraph("Second paragraph.")

    parsed = parse(save(document))

    assert parsed.text == "Annual report\n\nFirst paragraph.\n\nSecond paragraph."
    assert parsed.title is None
    assert parsed.metadata == {}


def test_tables_are_read_in_document_order() -> None:
    document = new_document()
    document.add_paragraph("Before the table.")
    table = document.add_table(rows=2, cols=2)
    for row, values in enumerate([["Name", "Score"], ["Ann", "12"]]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    document.add_paragraph("After the table.")

    parsed = parse(save(document))

    assert parsed.text.split("\n\n") == [
        "Before the table.",
        "Name",
        "Score",
        "Ann",
        "12",
        "After the table.",
    ]


def test_merged_cells_are_read_once() -> None:
    document = new_document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Wide"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Tall"
    table.cell(1, 1).text = "B"
    table.cell(2, 1).text = "C"

    parsed = parse(save(document))

    assert parsed.text.split("\n\n") == ["Wide", "Tall", "B", "C"]


def test_nested_tables() -> None:
    document = new_document()
    outer = document.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "Outer"
    inner = outer.cell(0, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Inner"

    assert parse(save(document)).text == "Outer\n\nInner"


def test_title_and_author_come_from_the_properties() -> None:
    document = new_document(title="  Quarterly   review ", author="Jane Doe")
    document.add_paragraph("Body")

    parsed = parse(save(document))

    assert parsed.title == "Quarterly review"
    assert parsed.metadata == {"author": "Jane Doe"}


def test_empty_document_is_allowed() -> None:
    parsed = parse(save(new_document()))

    assert parsed.text == ""
    assert parsed.title is None


def tiny_png() -> bytes:
    """A 1x1 grey PNG image."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        crc = zlib.crc32(kind + data)
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", crc)

    header = struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0)
    pixels = zlib.compress(b"\x00\x80")
    signature = b"\x89PNG\r\n\x1a\n"
    return signature + chunk(b"IHDR", header) + chunk(b"IDAT", pixels) + chunk(b"IEND", b"")


def test_images_are_ignored() -> None:
    document = new_document()
    document.add_paragraph("Caption above.")
    document.add_picture(io.BytesIO(tiny_png()))

    assert parse(save(document)).text == "Caption above."


@pytest.mark.parametrize("data", [b"", b"not a docx file", b"PK\x03\x04broken"])
def test_invalid_file_is_rejected(data: bytes) -> None:
    with pytest.raises(DocumentParsingError, match="Document is not a readable DOCX file."):
        parse(data)


def test_zip_that_is_not_a_docx_is_rejected() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "hello")

    with pytest.raises(DocumentParsingError, match="Document is not a readable DOCX file."):
        parse(buffer.getvalue())


def test_large_file_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    data = save(new_document())
    monkeypatch.setattr(docx_document, "MAX_DOCX_BYTES", len(data) - 1)

    with pytest.raises(DocumentParsingError, match="DOCX file is larger than"):
        parse(data)


def test_file_that_unpacks_to_too_much_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(docx_document, "MAX_UNPACKED_BYTES", 1000)

    with pytest.raises(DocumentParsingError, match="unpacks to too much data"):
        parse(save(new_document()))
