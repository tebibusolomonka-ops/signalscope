import io
import uuid
from pathlib import Path

import docx
import pytest
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from sample_files import make_pdf
from signalscope.domain.documents.asset import DocumentAsset
from signalscope.domain.documents.asset_service import DocumentAssetService
from signalscope.domain.documents.chunk import DocumentChunk
from signalscope.domain.documents.chunk_repository import DocumentChunkRepository
from signalscope.domain.documents.model import Document
from signalscope.domain.processing.processor import DocumentProcessor
from signalscope.domain.sources.model import Source, SourceType
from signalscope.parsing.docx_document import DOCX_CONTENT_TYPE
from signalscope.parsing.registry import create_default_parser_registry
from signalscope.storage.local import LocalBlobStore

pytestmark = pytest.mark.anyio

# Long enough that each page is split into several chunks.
PAGE_TEXT = [" ".join(f"page{page}word{word}" for word in range(150)) for page in (1, 2, 3)]


@pytest.fixture
def blobs(tmp_path: Path) -> LocalBlobStore:
    return LocalBlobStore(tmp_path / "blobs")


@pytest.fixture
async def source(session_factory: async_sessionmaker[AsyncSession]) -> Source:
    async with session_factory() as session:
        source = Source(type=SourceType.UPLOAD, name="Files")
        session.add(source)
        await session.commit()
    return source


async def process(
    session_factory: async_sessionmaker[AsyncSession],
    blobs: LocalBlobStore,
    source: Source,
    data: bytes,
    content_type: str,
) -> tuple[DocumentAsset, Document, list[DocumentChunk]]:
    async with session_factory() as session:
        document = Document(source_id=source.id, title="Report")
        session.add(document)
        await session.commit()
        asset = await DocumentAssetService(session, blobs).attach(
            document.id, filename="report", content_type=content_type, data=data
        )
    processor = DocumentProcessor(session_factory, blobs, create_default_parser_registry())
    await processor.process(asset.id)
    return (asset, *await stored(session_factory, asset.document_id))


async def stored(
    session_factory: async_sessionmaker[AsyncSession], document_id: uuid.UUID
) -> tuple[Document, list[DocumentChunk]]:
    async with session_factory() as session:
        document = await session.get(Document, document_id)
        chunks = await DocumentChunkRepository(session).list_by_document(document_id)
    assert document is not None
    return document, chunks


def check_offsets(document: Document, chunks: list[DocumentChunk]) -> None:
    assert document.content is not None
    for chunk in chunks:
        assert document.content[chunk.start_char : chunk.end_char] == chunk.text


async def test_pdf_chunks_know_their_page(
    session_factory: async_sessionmaker[AsyncSession], blobs: LocalBlobStore, source: Source
) -> None:
    _, document, chunks = await process(
        session_factory, blobs, source, make_pdf(PAGE_TEXT), "application/pdf"
    )

    assert len(chunks) > 3
    check_offsets(document, chunks)
    pages = [chunk.chunk_metadata["page_number"] for chunk in chunks]
    assert pages == sorted(pages)
    assert set(pages) == {1, 2, 3}
    for chunk in chunks:
        page = chunk.chunk_metadata["page_number"]
        assert chunk.chunk_metadata["section_kind"] == "page"
        assert chunk.chunk_metadata["section_index"] == page - 1
        # No chunk mixes words from two pages.
        assert all(word.startswith(f"page{page}word") for word in chunk.text.split())


async def test_docx_chunks_know_their_section(
    session_factory: async_sessionmaker[AsyncSession], blobs: LocalBlobStore, source: Source
) -> None:
    word_document = docx.Document()
    word_document.add_paragraph("Opening words.")
    word_document.add_heading("Findings", level=1)
    word_document.add_paragraph("What was found.")
    buffer = io.BytesIO()
    word_document.save(buffer)

    _, document, chunks = await process(
        session_factory, blobs, source, buffer.getvalue(), DOCX_CONTENT_TYPE
    )

    check_offsets(document, chunks)
    assert [(chunk.text, chunk.chunk_metadata) for chunk in chunks] == [
        ("Opening words.", {"section_kind": "section", "section_index": 0}),
        (
            "Findings\n\nWhat was found.",
            {"heading": "Findings", "section_kind": "section", "section_index": 1},
        ),
    ]


async def test_text_file_is_one_section(
    session_factory: async_sessionmaker[AsyncSession], blobs: LocalBlobStore, source: Source
) -> None:
    _, document, chunks = await process(
        session_factory, blobs, source, b"A short note.", "text/plain"
    )

    check_offsets(document, chunks)
    assert [chunk.chunk_metadata for chunk in chunks] == [
        {"section_kind": "document", "section_index": 0}
    ]


async def test_reprocessing_replaces_the_metadata(
    session_factory: async_sessionmaker[AsyncSession], blobs: LocalBlobStore, source: Source
) -> None:
    asset, _, first = await process(
        session_factory,
        blobs,
        source,
        make_pdf(["Old first page.", "Old second page."]),
        "application/pdf",
    )
    assert [chunk.chunk_metadata["page_number"] for chunk in first] == [1, 2]
    await blobs.put(asset.storage_key, make_pdf([None, None, "Only page three has text."]))

    processor = DocumentProcessor(session_factory, blobs, create_default_parser_registry())
    await processor.process(asset.id)

    document, chunks = await stored(session_factory, asset.document_id)
    check_offsets(document, chunks)
    assert [(chunk.text, chunk.chunk_metadata) for chunk in chunks] == [
        (
            "Only page three has text.",
            {"page_number": 3, "section_kind": "page", "section_index": 0},
        )
    ]
